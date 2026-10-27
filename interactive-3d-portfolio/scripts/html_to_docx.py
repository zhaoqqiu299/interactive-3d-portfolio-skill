# Convert a portfolio document (an HTML fragment with <aside class="dn"> design notes) to Word.
# Usage: python html_to_docx.py doc.html out.docx "Title" [site_root_for_images]
# Needs pandoc and python-docx.
import sys, re, subprocess, os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
src, out, title = sys.argv[1], sys.argv[2], sys.argv[3]
html = open(src).read()
html = re.sub(r'<aside class="dn">(.*?)</aside>', lambda m: '<div custom-style="Design Note"><p>' + m.group(1).replace('</b>', '</b> ', 1) + '</p></div>', html, flags=re.S)
html = re.sub(r'<span class="tag">(.*?)</span>', r'<em>[\1]</em>', html)
html = html.replace('<p class="sub">', '<p custom-style="Subtitle">')
tmp = out + '.tmp.html'
open(tmp, 'w').write(f'<html><head><meta charset="utf-8"><title>{title}</title></head><body>{html}</body></html>')
subprocess.run(['pandoc', tmp, '-f', 'html', '-t', 'docx', '-o', out, '--metadata', f'title={title}', '--resource-path=' + (sys.argv[4] if len(sys.argv) > 4 else '.')], check=True)
os.remove(tmp)
d = Document(out)
st = d.styles
def shade(p, fill):
    pPr = p._p.get_or_add_pPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill); pPr.append(shd)
    bdr = OxmlElement('w:pBdr')
    for side in ('top', 'bottom', 'left', 'right'):
        e = OxmlElement('w:' + side); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '6'); e.set(qn('w:color'), 'E3C58E'); bdr.append(e)
    pPr.append(bdr)
for p in d.paragraphs:
    if p.style.name == 'Design Note':
        shade(p, 'FBEFD2')
        for r in p.runs: r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x5B, 0x3A, 0x22)
        if p.runs: p.runs[0].font.color.rgb = RGBColor(0x9A, 0x5A, 0x14)
# the title paragraph pandoc adds from metadata duplicates the h1; drop it
for p in d.paragraphs:
    if p.style.name == 'Title': p._p.getparent().remove(p._p); break
for t in d.tables:
    tblPr = t._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:space'),'0'); e.set(qn('w:color'),'C8B48A'); b.append(e)
    tblPr.append(b)
d.save(out)
